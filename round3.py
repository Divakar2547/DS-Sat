from pathlib import Path
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

# ---------------- Paths (works from project root or output folder) ----------------
BASE = Path(__file__).parent
SRC = next(p for p in [BASE / "cleaned_network_flows.csv",
                       BASE / "output" / "cleaned_network_flows.csv"] if p.exists())
OUT = SRC.parent / "NetShield_Pitch.pptx"
NORMAL = "normal"

NAVY = RGBColor(0x1F, 0x38, 0x64)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xEE, 0xF2, 0xF8)
DARK = RGBColor(0x22, 0x22, 0x22)

# ---------------- Data ----------------
df = pd.read_csv(SRC).replace([np.inf, -np.inf], np.nan)
df["label"] = df["label"].astype(str).str.strip().str.lower()
df["unusual"] = (df["label"] != NORMAL).astype(int)
total, un = len(df), int(df["unusual"].sum())
overall = un / total

def rate_table(col):
    return df.groupby(col)["unusual"].agg(rate="mean", unusual="sum", total="count")

def top_of(g, min_n=20):
    gg = g[g["total"] >= min_n]
    if gg.empty:
        gg = g
    idx = gg["rate"].idxmax()
    return idx, gg.loc[idx]

def lift(rate):
    return rate / overall if overall else 0

def seg_rate(col, value):
    if col not in df.columns:
        return float("nan")
    sub = df[df[col] == value]
    return sub["unusual"].mean() if len(sub) else float("nan")

# ---------------- Slide helpers ----------------
prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def new_slide(title_text, notes=""):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.2), Inches(12), Inches(0.8))
    p = tb.text_frame.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = WHITE
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s

def add_bullets(slide, items, x, y, w, h, size=18):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "\u2022 " + t
        p.font.size = Pt(size)
        p.space_after = Pt(10)

def add_box(slide, heading, items, x, y, w, h, size=14):
    box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT
    box.line.color.rgb = NAVY
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = heading
    p.font.size = Pt(size + 2)
    p.font.bold = True
    p.font.color.rgb = NAVY
    for t in items:
        q = tf.add_paragraph()
        q.text = t
        q.font.size = Pt(size)
        q.font.color.rgb = DARK
        q.space_before = Pt(6)

def chart_slide(title_text, cats, vals, evidence, meaning, fmt="0%", notes="", label_size=11):
    s = new_slide(title_text, notes)
    cd = CategoryChartData()
    cd.categories = [str(c) for c in cats]
    cd.add_series("Series", [float(v) for v in vals])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.4),
                            Inches(7.8), Inches(5.6), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    plot = ch.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format = fmt
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.size = Pt(label_size)
    plot.series[0].format.fill.solid()
    plot.series[0].format.fill.fore_color.rgb = NAVY
    ch.value_axis.has_major_gridlines = False
    ch.value_axis.tick_labels.number_format = fmt
    ch.value_axis.tick_labels.number_format_is_linked = False
    ch.value_axis.tick_labels.font.size = Pt(11)
    ch.category_axis.tick_labels.font.size = Pt(11)
    add_box(s, "Evidence", evidence, 8.6, 1.4, 4.3, 2.9)
    add_box(s, "Business meaning", meaning, 8.6, 4.45, 4.3, 2.5)
    return s

# ---------------- Compute insights ----------------
g_proto = rate_table("protocol")
g_hour = rate_table("hour").sort_index()
g_src = rate_table("source_zone")
g_port = rate_table("port_category")
g_group = rate_table("host_group")
g_dst = rate_table("destination_zone")
g_pair = df.groupby(["source_zone", "destination_zone"])["unusual"].agg(
    rate="mean", unusual="sum", total="count")

p_name, p_row = top_of(g_proto)
h_name, h_row = top_of(g_hour, 10)
s_name, s_row = top_of(g_src)
d_name, d_row = top_of(g_dst)
pt_name, pt_row = top_of(g_port)
hg_name, hg_row = top_of(g_group)
pair_name, pair_row = top_of(g_pair, 15)

off_rate = seg_rate("is_off_hours", 1)
biz_rate = seg_rate("is_off_hours", 0)
mw_rate = seg_rate("maintenance_window", 1)
nmw_rate = seg_rate("maintenance_window", 0)
wk_rate = seg_rate("is_weekend", 1)

feat = ["bytes", "packets", "duration", "bytes_per_packet", "bytes_per_second"]
med = df.groupby("unusual")[feat].median().T
mean_ = df.groupby("unusual")[feat].mean().T
if 0 not in med.columns or 1 not in med.columns:
    raise SystemExit("Need both normal and unusual flows to compare features.")
med.columns = ["normal", "unusual"]
mean_.columns = ["normal", "unusual"]
med["ratio"] = med["unusual"] / med["normal"].replace(0, np.nan)
big_feat = med["ratio"].sub(1).abs().idxmax()
big_ratio = med.loc[big_feat, "ratio"]
direction = "higher" if big_ratio >= 1 else "lower"

events = df.loc[df["unusual"] == 1, "label"].value_counts()
extreme = int(df["extreme_bytes_flag"].sum()) if "extreme_bytes_flag" in df.columns else 0

# ---------------- Slide 1: title ----------------
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
tb = s.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.5), Inches(3))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Defensive Monitoring Priorities"
p.font.size = Pt(46)
p.font.bold = True
p.font.color.rgb = WHITE
p = tf.add_paragraph()
p.text = "What separates normal from unusual network traffic"
p.font.size = Pt(24)
p.font.color.rgb = WHITE
p = tf.add_paragraph()
p.text = "NetShield Security Lab  |  Cyber Traffic Insight"
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(0xBB, 0xC8, 0xE0)
s.notes_slide.notes_text_frame.text = (
    "Introduce the question: which traffic characteristics repeatedly mark unusual "
    "activity, and where should defenders focus first?")

# ---------------- Slide 2: context ----------------
s = new_slide("Context and approach",
              "Explain scope: anonymized flow data, defensive analysis only, no offensive testing.")
add_bullets(s, [
    f"{total:,} anonymized network flows analysed from five linked files "
    "(flows, hosts, sessions, labels, calendar)",
    f"{un:,} flows ({overall:.1%}) carry an event label; the rest are labelled normal",
    "Data quality fixed first: missing protocols, duplicate flow IDs, inconsistent labels, "
    f"extreme byte counts ({extreme} flagged), unmatched host groups",
    "Features engineered: bytes per packet, bytes per second, packets per second, "
    "log bytes, off-hours and weekend flags",
    "Scope is defensive: understand patterns to prioritise monitoring, no offensive testing",
], 0.7, 1.5, 12, 5.5, size=20)

# ---------------- Slide 3: executive summary ----------------
s = new_slide("Executive summary",
              "Give the headline findings in under a minute before diving into evidence.")
add_bullets(s, [
    f"Overall unusual rate is {overall:.1%}; several segments run well above this baseline",
    f"Highest-risk protocol: {p_name} at {p_row['rate']:.1%} ({lift(p_row['rate']):.1f}x baseline)",
    f"Highest-risk port category: {pt_name} ({pt_row['rate']:.1%}); "
    f"host group: {hg_name} ({hg_row['rate']:.1%})",
    f"Peak unusual hour: {int(h_name):02d}:00 at {h_row['rate']:.1%}",
    f"Biggest feature gap: median {big_feat} is {big_ratio:.2f}x of normal traffic ({direction})",
    "Recommendation: tune monitoring on the recurring segments below, not uniformly",
], 0.7, 1.5, 12, 5.5, size=20)

# ---------------- Slide 4: protocol ----------------
chart_slide("Protocol distribution and risk",
            g_proto.index, g_proto["rate"],
            [f"{r.Index}: {int(r.total):,} flows, {r.rate:.1%} unusual" for r in g_proto.itertuples()],
            [f"{p_name} is the most alert-prone protocol ({lift(p_row['rate']):.1f}x baseline)",
             "Protocol-specific alert rules give the best return on tuning effort"],
            notes="Show which protocols carry volume and which carry risk. They are not always the same.")

# ---------------- Slide 5: feature comparison ----------------
s = new_slide("Feature comparison: normal vs unusual traffic",
              "Medians are used because byte counts are skewed by extreme values.")
rows = len(feat) + 1
tbl = s.shapes.add_table(rows, 5, Inches(0.6), Inches(1.5), Inches(12.1), Inches(3.2)).table
for j, h in enumerate(["Feature", "Median normal", "Median unusual", "Mean unusual", "Median ratio"]):
    tbl.cell(0, j).text = h
for i, f in enumerate(feat, start=1):
    vals = [f, f"{med.loc[f, 'normal']:,.2f}", f"{med.loc[f, 'unusual']:,.2f}",
            f"{mean_.loc[f, 'unusual']:,.2f}", f"{med.loc[f, 'ratio']:.2f}x"]
    for j, v in enumerate(vals):
        tbl.cell(i, j).text = v
for r in range(rows):
    for c in range(5):
        for p in tbl.cell(r, c).text_frame.paragraphs:
            p.font.size = Pt(16)
add_bullets(s, [
    f"Largest difference: {big_feat} (median ratio {big_ratio:.2f}x)",
    "Unusual flows differ in size and rate characteristics, so thresholds on these "
    "features can support detection",
    "Medians are compared, not means, to avoid distortion by extreme byte counts",
], 0.7, 5.0, 12, 2.2, size=17)

# ---------------- Slide 6: hourly ----------------
chart_slide("Hourly traffic pattern",
            [f"{int(h):02d}" for h in g_hour.index], g_hour["rate"],
            [f"Peak unusual hour: {int(h_name):02d}:00 ({h_row['rate']:.1%})",
             f"Off-hours unusual rate: {off_rate:.1%} vs business hours {biz_rate:.1%}",
             f"Maintenance window: {mw_rate:.1%} vs outside {nmw_rate:.1%}",
             f"Weekend unusual rate: {wk_rate:.1%}"],
            ["Analyst coverage and alert sensitivity should follow the risk curve by hour",
             "Maintenance windows need their own baseline to reduce false positives"],
            label_size=8,
            notes="Link the hourly curve to staffing and alert tuning.")

# ---------------- Slide 7: zones ----------------
chart_slide("Zone comparison (source zone)",
            g_src.index, g_src["rate"],
            [f"Highest source zone: {s_name} ({s_row['rate']:.1%})",
             f"Highest destination zone: {d_name} ({d_row['rate']:.1%})",
             f"Riskiest path: {pair_name[0]} to {pair_name[1]} "
             f"({pair_row['rate']:.1%}, {int(pair_row['total'])} flows)"],
            ["Zone boundaries are natural places to enforce and monitor policy",
             "Prioritise the riskiest zone paths for tighter flow logging"],
            notes="Use the ZoneMatrix sheet in the Excel workbook as supporting evidence.")

# ---------------- Slide 8: ports ----------------
chart_slide("Port category analysis",
            g_port.index, g_port["rate"],
            [f"{r.Index}: {int(r.total):,} flows, {r.rate:.1%} unusual" for r in g_port.itertuples()],
            [f"{pt_name} traffic is the most alert-prone ({lift(pt_row['rate']):.1f}x baseline)",
             "Port-category rules narrow the monitoring scope instead of watching everything equally"],
            notes="Port categories show where unusual traffic concentrates.")

# ---------------- Slide 9: host groups ----------------
chart_slide("Host-group alert summary",
            g_group.index, g_group["rate"],
            [f"{r.Index}: {int(r.unusual)} unusual of {int(r.total):,} ({r.rate:.1%})"
             for r in g_group.itertuples()],
            [f"{hg_name} has the highest alert rate and should be reviewed first",
             "Group-specific baselines stop noisy groups from hiding real changes"],
            notes="Host groups are where ownership lies, so this is where management actions land.")

# ---------------- Slide 10: event categories ----------------
if len(events) > 0:
    chart_slide("Unusual traffic by event category",
                list(events.index), list(events.values),
                [f"{k}: {v:,} flows ({v / un:.0%} of unusual)" for k, v in events.items()],
                ["Different event types need different response playbooks",
                 "Track the largest categories as monitoring KPIs"],
                fmt="#,##0",
                notes="Labels were standardised in cleaning so categories are not split by spelling.")

# ---------------- Slide 11: management attention ----------------
s = new_slide("Areas needing management attention",
              "Prioritise by recurring risk, not by one-off spikes.")
add_bullets(s, [
    f"{p_name} traffic: highest alert rate among protocols ({p_row['rate']:.1%})",
    f"{pt_name} port category: {pt_row['rate']:.1%} unusual, {lift(pt_row['rate']):.1f}x baseline",
    f"{hg_name} host group: top alert rate at {hg_row['rate']:.1%}",
    f"Zone path {pair_name[0]} to {pair_name[1]}: {pair_row['rate']:.1%} unusual",
    f"Hour {int(h_name):02d}:00 and off-hours traffic: elevated risk window",
    f"{extreme} extreme-byte flows need a standing review process",
], 0.7, 1.5, 12, 5.5, size=20)

# ---------------- Slide 12: recommendations ----------------
s = new_slide("Recommendations: defensive monitoring priorities",
              "Each recommendation links to a finding and can be owned by the security team.")
add_bullets(s, [
    f"1. Create dedicated alert rules for {p_name} traffic and the {pt_name} port category",
    f"2. Raise logging and review depth on the {pair_name[0]} to {pair_name[1]} zone path",
    f"3. Start host-group reviews with {hg_name}, then rank the others by alert rate",
    f"4. Align analyst coverage with the peak risk hour ({int(h_name):02d}:00) and off-hours",
    "5. Use separate baselines for maintenance windows to reduce false positives",
    f"6. Monitor threshold shifts in {big_feat}, which shows the clearest gap",
    "7. Review weekly: unusual rate by protocol, port, zone and host group on one dashboard",
], 0.7, 1.5, 12, 5.6, size=19)

# ---------------- Slide 13: close ----------------
s = new_slide("Summary and next steps",
              "Close by restating the ask: approve the monitoring priorities.")
add_bullets(s, [
    f"Unusual traffic ({overall:.1%} of flows) concentrates in a few recurring segments",
    "Focusing monitoring there gives more coverage per analyst hour than uniform alerting",
    "Next steps: pilot the alert rules, track the unusual rate per segment, review monthly",
    "Limitations: labels come from historical classification; thresholds should be validated on new data",
], 0.7, 1.5, 12, 5, size=20)

prs.save(str(OUT))
print("Saved", OUT)