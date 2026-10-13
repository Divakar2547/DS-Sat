from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

BASE = Path(__file__).parent
SRC = BASE / "cleaned_flows.csv"
if not SRC.exists():
    SRC = BASE / "output" / "cleaned_network_flows.csv"
if not SRC.exists():
    raise FileNotFoundError(
        f"Could not find cleaned_flows.csv or output/cleaned_network_flows.csv under {BASE}"
    )

df = pd.read_csv(SRC)
df["unusual"] = (df["label"].str.lower() != "normal").astype(int)

def rate_by(col):
    g = df.groupby(col)["unusual"].agg(["mean", "sum", "count"]).sort_values("mean", ascending=False)
    return g

prs = Presentation(); prs.slide_width = Inches(13.33); prs.slide_height = Inches(7.5)

def text_slide(title, bullets):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    tf = s.placeholders[1].text_frame; tf.clear()
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b; p.font.size = Pt(20)

def chart_slide(title, series, takeaway):
    s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = title
    cd = CategoryChartData(); cd.categories = [str(i) for i in series.index]
    cd.add_series("Unusual rate", [round(float(v), 4) for v in series.values])
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(1.5),
                       Inches(8.5), Inches(5.2), cd)
    tb = s.shapes.add_textbox(Inches(9.4), Inches(1.8), Inches(3.5), Inches(4)).text_frame
    tb.word_wrap = True; tb.text = takeaway; tb.paragraphs[0].font.size = Pt(18)

total, un = len(df), df["unusual"].sum()
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Defensive Monitoring Priorities"
s.placeholders[1].text = "NetShield Security Lab | Network-flow insight"

text_slide("Context & approach", [
    f"{total:,} anonymized flows analysed; {un:,} ({un/total:.1%}) labelled unusual",
    "Data cleaned: missing protocols, duplicate IDs, inconsistent labels, extreme bytes",
    "Question: what separates normal from unusual traffic?",
    "Scope: defensive monitoring only"])

for col, name in [("protocol", "Protocol"), ("hour", "Hour"), ("port_category", "Port category"),
                  ("host_group", "Host group")]:
    g = rate_by(col)
    top = g.index[0]
    chart_slide(f"Unusual rate by {name}", g["mean"].head(8),
                f"{top} has the highest unusual rate ({g['mean'].iloc[0]:.1%}) "
                f"vs overall {un/total:.1%}.\n\nBusiness meaning: [write this]")

text_slide("Recommendations", [
    "1. Add alert rules for the highest-risk protocol / port combinations",
    "2. Increase monitoring during peak unusual hours",
    "3. Prioritise review of top host groups by alert rate",
    "4. Baseline per-zone traffic and alert on deviations",
    "5. Review maintenance-window traffic separately to cut false positives"])
OUT = BASE / "NetShield_Pitch.pptx"
prs.save(OUT); print("saved", OUT)